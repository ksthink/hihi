#!/usr/bin/env python3
"""산림청 산코드 두 소스를 병합해 data/mnt-codes.json 을 생성한다.

소스 (scripts/):
  1) MNT_CODE.xlsx — 산코드 목록 (순번/산이름/위치/산코드, 2,931건)
  2) mnt.xlsx      — 산림청(항공본부) 산정보 (산코드/산명/소재지/높이 등 13열, 4,704건)

병합 정책 (코드 = 유일 키):
  - 두 소스의 합집합. 공통 코드는 산정보(mnt.xlsx)가 우선 —
    이름 불일치 12건(예: 아가봉→아기봉)은 산정보 쪽이 정본.
  - region: 산정보 소재지(리 단위까지 상세) 우선, 없으면 MNT_CODE 위치.
  - elev: 산정보 높이(m), 0/누락은 제외.

산코드(9자리)가 전 시스템의 mountain_id 표준이다:
  Storage packs/<산코드>/ · mountains.id · saved_packs/climb_records.mountain_id ·
  app.js PARKS/MNT 키 · IndexedDB 팩 키

출력: [{"code","name","region","elev?"}, ...]  (이름은 전국 중복 → 표시 시 region 병기)

사용: (venv 에 openpyxl 필요)
  python3 scripts/convert_mnt_codes.py
"""
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl 필요: pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_CODE = os.path.join(ROOT, "scripts", "MNT_CODE.xlsx")
SRC_INFO = os.path.join(ROOT, "scripts", "mnt.xlsx")
DST = os.path.join(ROOT, "data", "mnt-codes.json")

norm = lambda v: " ".join(str(v or "").split())

# 1) 코드 목록 (순번, 산이름, 위치, 산코드)
merged = {}
wb = openpyxl.load_workbook(SRC_CODE, read_only=True)
for _, name, region, code in list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:]:
    if not code or not name:
        continue
    code = str(code).strip()
    merged[code] = {"code": code, "name": norm(name), "region": norm(region)}

n_code_only = len(merged)

# 2) 산정보 (산코드, 산명, ..., 소재지[6], ..., 높이[11]) — 공통 코드는 이쪽이 우선
wb = openpyxl.load_workbook(SRC_INFO, read_only=True)
added = updated = 0
for r in list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:]:
    if r[0] is None or not r[1]:
        continue
    code = str(r[0]).strip()
    entry = merged.get(code)
    if entry is None:
        entry = {"code": code, "name": "", "region": ""}
        merged[code] = entry
        added += 1
    else:
        updated += 1
    entry["name"] = norm(r[1]) or entry["name"]
    entry["region"] = norm(r[6]) or entry["region"]
    try:
        elev = float(r[11]) if r[11] is not None else 0
    except (TypeError, ValueError):
        elev = 0
    if elev > 0:
        entry["elev"] = round(elev, 1)

out = sorted(merged.values(), key=lambda m: m["code"])
with open(DST, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))

print(f"MNT_CODE {n_code_only}건 + 산정보 병합(갱신 {updated}·신규 {added})"
      f" → 총 {len(out)}건, {os.path.getsize(DST)//1024}KB → {DST}")
