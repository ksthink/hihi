#!/usr/bin/env python3
"""하이하잇 관리자 서버 — 코스 큐레이션·팩 배포 콘솔.

serve.py(정적+PMTiles 프록시)를 상속하고 /api/* JSON 라우트와 백그라운드 잡을 얹는다.
사용자 앱(/)과 관리자(/admin/)를 같은 포트에서 서빙하므로 serve.py 대체로 쓰면 된다.

운영자 1인용. 원격(EC2 등) 접속을 위해 기본 0.0.0.0 바인딩이며, /api 는
**토큰 인증**(.env ADMIN_TOKEN, 최초 실행 시 자동 생성)으로 보호한다.
localhost(127.0.0.1) 요청은 토큰 없이 허용. Supabase secret 키는 .env 로만
로드(클라이언트로 절대 전달 안 함). 로컬 전용으로 잠그려면 ADMIN_BIND=127.0.0.1.

사용: .venv/bin/python scripts/admin_server.py [port]   (기본 8890)
접속: http://<호스트>:8890/admin/?token=<ADMIN_TOKEN>   (토큰은 시작 로그에 출력)
"""
import json
import os
import queue
import re
import secrets
import subprocess
import sys
import threading
import time
import traceback
import urllib.parse
import urllib.request
import uuid
from functools import partial
from http.server import ThreadingHTTPServer

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from serve import Handler as BaseHandler, ROOT  # noqa: E402

import draft_store  # noqa: E402
import pack_lib as pl  # noqa: E402
import sheet_lib  # noqa: E402

PORT = int(sys.argv[1]) if len(sys.argv) > 1 else 8890
BIND = os.environ.get("ADMIN_BIND", "0.0.0.0")

# 표시 정책 기본값 — 값 = {zoom: 표시 시작 줌(0=항상, null=끔), icon: 기호, size: 글자 px,
# bold: 볼드}. 레거시 파일/요청(값이 숫자|null)은 _norm_display_cat 이 zoom 으로 승격.
# 앱(app.js SPOT_DISPLAY_DEFAULT · basemap-style.js POI_DISPLAY_DEFAULT)과 일치 유지.
SPOT_DISPLAY_DEFAULT = {"version": 2, "categories": {
    "정상": {"zoom": 0, "icon": True, "size": 14, "bold": True},
    "장소": {"zoom": 14, "icon": True, "size": 8, "bold": False},
    "조망점": {"zoom": 18, "icon": True, "size": 8, "bold": False},
    "화장실": {"zoom": 18, "icon": True, "size": 8, "bold": False},
    "정자": {"zoom": 18, "icon": True, "size": 8, "bold": False},
    "헬기장": {"zoom": 18, "icon": True, "size": 8, "bold": False},
    "음수대": {"zoom": 18, "icon": True, "size": 8, "bold": False},
    "주차장": {"zoom": None, "icon": True, "size": 8, "bold": False},
    "분기점": {"zoom": None, "icon": True, "size": 8, "bold": False},
    "시종점": {"zoom": None, "icon": True, "size": 8, "bold": False},
}}

POI_DISPLAY_DEFAULT = {"version": 2, "categories": {
    "전철역": {"zoom": 12, "icon": True, "size": 10, "bold": False},
    "버스정류장": {"zoom": 14, "icon": True, "size": 8, "bold": False},
    "사찰": {"zoom": 14, "icon": True, "size": 10, "bold": False},
    "편의시설": {"zoom": 14, "icon": True, "size": 8, "bold": False},
    "학교": {"zoom": 14, "icon": False, "size": 10, "bold": False},
    "관공서": {"zoom": 14, "icon": False, "size": 10, "bold": False},
    "병원": {"zoom": 14, "icon": False, "size": 10, "bold": False},
    "아파트단지": {"zoom": 14, "icon": False, "size": 10, "bold": False},
    "공원": {"zoom": 14, "icon": False, "size": 10, "bold": False},
    "마트·쇼핑": {"zoom": 15, "icon": False, "size": 10, "bold": False},
    "문화·체육": {"zoom": 15, "icon": False, "size": 10, "bold": False},
}}


# ── 배포 버전 (관리자 콘솔 우상단 표시) ──
# EC2 등 원격 체크아웃에서 "git pull 이 반영됐는지"를 눈으로 확인하기 위한 것.
# ⚠️ 시작 시점이 아니라 **요청 시점**에 읽는다 — 정적 파일(admin/*.js·css)만 pull 한
#    경우엔 서버를 재시작하지 않으므로, 시작 시 캐시하면 옛 커밋을 계속 보여준다.
_VER_CACHE = {"t": 0.0, "v": None}


def git_version():
    """{commit, date, subject, dirty} — git 이 없거나 실패하면 빈 값."""
    now = time.time()
    if _VER_CACHE["v"] and now - _VER_CACHE["t"] < 5:   # 연타 방지용 짧은 캐시
        return _VER_CACHE["v"]

    def git(*args):
        try:
            r = subprocess.run(["git", "-C", ROOT, *args],
                               capture_output=True, text=True, timeout=3)
            return r.stdout.strip() if r.returncode == 0 else ""
        except Exception:
            return ""

    v = {"commit": git("rev-parse", "--short", "HEAD"),
         "date": git("log", "-1", "--format=%cd", "--date=format:%Y-%m-%d %H:%M"),
         "subject": git("log", "-1", "--format=%s"),
         # admin_data/ 가 추적 대상이라 관리자 작업 후엔 dirty 가 뜬다 — pull 충돌 예고.
         "dirty": bool(git("status", "--porcelain"))}
    _VER_CACHE.update(t=now, v=v)
    return v


def _norm_display_cat(cat, v, default):
    """카테고리 설정값 정규화·검증. 레거시(숫자|null)는 zoom 으로 승격."""
    if v is None or isinstance(v, (int, float)):
        v = {"zoom": v}
    if not isinstance(v, dict):
        raise ValueError(f"{cat}: 숫자·null 또는 객체 필요")
    out = {**default}
    for key, val in v.items():
        if key == "zoom":
            if val is not None and not (isinstance(val, (int, float)) and 0 <= val <= 22):
                raise ValueError(f"{cat}.zoom: 0~22 또는 null(끔)")
        elif key == "bold":
            if not isinstance(val, bool):
                raise ValueError(f"{cat}.bold: true/false 필요")
        elif key == "icon":
            # 3상: false=끔 · true=기본 아이콘 · 문자열=그 문자를 이름 앞에.
            # ⚠️ BMP(U+0000~FFFF) 만 허용한다 — 자체 호스팅 글리프 PBF 가 65535 에서
            #    끝나므로 이모지는 웹·네이티브 모두 아무 경고 없이 사라진다. 화면에서
            #    막고 있지만, 서버가 마지막 방어선이다.
            if isinstance(val, str):
                val = val.strip()
                if len(val) > 2:
                    raise ValueError(f"{cat}.icon: 기호는 2글자까지")
                bad = next((c for c in val if ord(c) > 0xFFFF), None)
                if bad:
                    raise ValueError(
                        f"{cat}.icon: '{bad}' 는 글리프가 없어 지도에 안 그려집니다 "
                        "(이모지 불가 — ★ ▲ ● ■ ◆ ⊕ ✚ ⚑ ㉿ ㈜ 등 BMP 기호를 쓰세요)")
                if not val:
                    val = True   # 빈 문자열 = 기본 아이콘
            elif not isinstance(val, bool):
                raise ValueError(f"{cat}.icon: true/false 또는 기호 문자")
        elif key == "size":
            if not (isinstance(val, (int, float)) and 6 <= val <= 24):
                raise ValueError(f"{cat}.size: 6~24 px")
        else:
            raise ValueError(f"{cat}: 알 수 없는 속성 {key}")
        out[key] = val
    return out


def _norm_display_cfg(cats, defaults):
    """categories 전체 정규화 — 누락 분류는 기본값, 알 수 없는 분류는 거부."""
    for k in cats:
        if k not in defaults["categories"]:
            raise ValueError(f"알 수 없는 분류: {k}")
    return {"version": 2, "categories": {
        k: _norm_display_cat(k, cats.get(k, d), d)
        for k, d in defaults["categories"].items()}}


def load_env():
    p = os.path.join(ROOT, ".env")
    if not os.path.exists(p):
        return
    for line in open(p, encoding="utf-8"):
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip())


def ensure_admin_token():
    """ADMIN_TOKEN 이 없으면 생성해 .env 에 영속화."""
    if os.environ.get("ADMIN_TOKEN"):
        return os.environ["ADMIN_TOKEN"]
    tok = secrets.token_hex(16)
    with open(os.path.join(ROOT, ".env"), "a", encoding="utf-8") as f:
        f.write(f"\n# 관리자 콘솔 /api 인증 토큰 (admin_server.py 자동 생성)\nADMIN_TOKEN={tok}\n")
    os.environ["ADMIN_TOKEN"] = tok
    return tok


# ── 백그라운드 잡 (직렬 1개 — 운영자 1인) ──
class JobManager:
    def __init__(self):
        self.jobs = {}
        self.q = queue.Queue()
        threading.Thread(target=self._worker, daemon=True).start()

    def submit(self, title, fn, code=None):
        jid = uuid.uuid4().hex[:12]
        job = {"id": jid, "title": title, "code": code, "state": "queued",
               "step": "대기 중", "progress": 0.0, "log": [], "result": None, "error": None}
        self.jobs[jid] = job
        self.q.put((job, fn))
        return jid

    def busy(self, code):
        """해당 산코드의 잡이 대기/실행 중인가 (삭제와 배포의 레이스 방지용)."""
        return any(j.get("code") == code and j["state"] in ("queued", "running")
                   for j in self.jobs.values())

    def _worker(self):
        while True:
            job, fn = self.q.get()
            job["state"] = "running"
            try:
                job["result"] = fn(job)
                job["state"] = "done"
                job["progress"] = 1.0
                job["step"] = "완료"
            except Exception as e:
                job["state"] = "error"
                job["error"] = str(e)
                job["log"].append(traceback.format_exc(limit=3))


JOBS = JobManager()


def job_step(job, step, progress):
    job["step"] = step
    job["progress"] = progress
    job["log"].append(step)


# ── 초안 생성/자동 시드 ──
def create_draft(code):
    """초안 생성 — 산림청 원본 기반 (정상 자동 시드 + DEM 프리페치).
    레거시 역임포트는 폐기: KNPS 탐방로 유래 코스가 점 간격 ~104m 로 거칠어
    (산림청 6.7m) 삭제한 산이 거친 선형으로 부활하는 통로였다.
    반환: (draft, job_id|None)."""
    if draft_store.load(code):
        raise ValueError(f"{code} 초안이 이미 있음")
    meta = next((m for m in _mnt_codes() if m["code"] == code), {})
    draft = draft_store.new_draft(code, meta)
    if meta.get("top100"):  # 100대 명산은 추천 노출(famous) 자동 체크
        draft["mountain"]["famous"] = True
    draft_store.save(code, draft)

    # 자동 코스 시드는 폐기 (파편화된 국립공원 구간망에서 괴물 코스 생성 — 부록 D).
    # 코스는 운영자가 GPX 업로드로 직접 입력. DEM 만 미리 받아 첫 recompute 지연을 줄인다.
    def seed(job):
        import dem_cache
        job_step(job, "DEM 타일 확보", 0.3)
        dem_cache.ensure(draft["mountain"]["bbox"],
                         log=lambda m: job["log"].append(m))
        job_step(job, "DEM 준비 완료 — 코스는 GPX 업로드로 입력", 0.95)
        return {"code": code}

    return draft, JOBS.submit(f"{draft['mountain']['name']} DEM 준비", seed, code=code)


def _next_manual_code():
    """수동 등록 산의 산코드 — '9'+8자리 예약 대역.
    산림청 코드는 전부 1~4로 시작하므로 9-접두는 절대 충돌하지 않는다.
    admin_data 의 기존 9######## 중 최대+1, 없으면 900000001."""
    base = draft_store.ADMIN_DATA
    used = [int(c) for c in (os.listdir(base) if os.path.isdir(base) else [])
            if re.fullmatch(r"9\d{8}", c)]
    return str((max(used) + 1) if used else 900000001)


_MNT_CACHE = None
_TOP100_CACHE = None


def _top100():
    """cache/top100.json (fetch_top100.py 산출) → {로컬 산코드: 항목}.
    100대 명산 분류(top100)와 정상 좌표·해발고도 소스. 파일 없으면 빈 dict."""
    global _TOP100_CACHE
    if _TOP100_CACHE is None:
        p = os.path.join(ROOT, "cache", "top100.json")
        rows = json.load(open(p, encoding="utf-8")) if os.path.exists(p) else []
        _TOP100_CACHE = {t["local_code"]: t for t in rows if t.get("local_code")}
    return _TOP100_CACHE


def _xlsx_meta():
    """산림청 xlsx(산 소개 시드용으로 보관 중) → {산코드: {region, elev}}.
    이름 중복(전국 317건, 예: 청계산 과천·가평·양평) 구분 표기에 쓴다.
    소재지·높이는 mnt.xlsx 우선, 없으면 MNT_CODE.xlsx 위치로 보충."""
    try:
        import openpyxl
    except ImportError:
        return {}
    meta = {}
    p = os.path.join(ROOT, "scripts", "MNT_CODE.xlsx")  # 순번/산이름/위치/산코드
    if os.path.exists(p):
        ws = openpyxl.load_workbook(p, read_only=True).active
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if r and r[3]:
                loc = str(r[2] or "").strip() or None
                meta[str(r[3]).strip()] = {"region": loc, "elev": None}
    p = os.path.join(ROOT, "scripts", "mnt.xlsx")  # 산코드/산명/…/소재지[6]/…/높이[11]
    if os.path.exists(p):
        ws = openpyxl.load_workbook(p, read_only=True).active
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or not r[0]:
                continue
            cur = meta.setdefault(str(r[0]).strip(), {"region": None, "elev": None})
            loc = str(r[6] or "").strip()
            if loc:
                cur["region"] = loc
            try:
                if r[11] and float(r[11]) > 0:
                    cur["elev"] = round(float(r[11]))
            except (TypeError, ValueError):
                pass
    return meta


def _mnt_codes():
    """mountain/<산코드>/ 스캔 — 내려받은 산림청 원본이 곧 산 목록.
    이름은 PMNTN_<이름>_<산코드>.json 파일명에서 추출, 지역·높이는 xlsx 메타로 보강."""
    global _MNT_CACHE
    if _MNT_CACHE is None:
        base = os.path.join(ROOT, "mountain")
        xm = _xlsx_meta()
        rows = []
        for code in sorted(os.listdir(base)) if os.path.isdir(base) else []:
            dp = os.path.join(base, code)
            if not (re.fullmatch(r"\d{9}", code) and os.path.isdir(dp)):
                continue
            name = code
            for fn in os.listdir(dp):
                m = re.fullmatch(rf"PMNTN_(.+)_{code}\.json", fn)
                if m and not m.group(1).startswith(("SPOT_", "SAFE_")):
                    name = m.group(1).replace("_", " ")
                    break
            m = xm.get(code) or {}
            t = _top100().get(code)  # 100대 명산 분류 + 정상 좌표·고도 보강
            rows.append({"code": code, "name": name,
                         "region": m.get("region") or (t and t["addr"]) or None,
                         # 해발고도: 기존 xlsx 값 우선, 없으면 100대 API 값으로 보충
                         "elev": m.get("elev") or (t and t["elev"] and round(t["elev"])) or None,
                         "top100": bool(t),
                         "peak": t and {"lat": t["lat"], "lon": t["lon"], "elev": t["elev"]}})
        _MNT_CACHE = rows
    return _MNT_CACHE


# 로그인 실패 잠금 (IP 별): 5회 실패 → 10분 잠금
LOGIN_FAILS = {}          # ip -> {"n": 실패 횟수, "until": 잠금 해제 시각}
MAX_FAILS, LOCK_SEC = 5, 600


# ── 등반 기록 (진단 열람) ──
# climb_records 의 RLS 는 "본인 기록만"이라 브라우저의 publishable 키로는 남의 기록을 못 읽는다.
# 관리자 콘솔 대신 이 서버가 service_role 키로 조회한다 — 키가 브라우저에 나가지 않고,
# 사용자 트랙 열람 권한을 상시 열어두는 RLS 정책도 필요 없다(2026-07-29 결정).
SUPABASE_URL = os.environ.get(
    "SUPABASE_URL", "https://durnojryhhsajnlwvdzt.supabase.co").rstrip("/")


def _supabase_secret():
    return (os.environ.get("SUPABASE_SECRET_KEY")
            or os.environ.get("SUPABASE_SERVICE_ROLE") or "")


def _records(limit=200):
    """등반 기록 목록 + 진단 요약. track.points 는 크므로 개수만 돌려준다."""
    key = _supabase_secret()
    if not key:
        raise RuntimeError(
            "SUPABASE_SECRET_KEY 가 없습니다 — .env 에 service_role 키를 넣고 서버를 재시작하세요")
    qs = urllib.parse.urlencode({
        "select": ("id,user_id,mountain_id,course_name,started_at,ended_at,"
                   "distance_km,ascent_m,duration_s,track"),
        "order": "started_at.desc",
        "limit": str(limit),
    })
    req = urllib.request.Request(
        f"{SUPABASE_URL}/rest/v1/climb_records?{qs}",
        headers={"apikey": key, "Authorization": f"Bearer {key}"})
    with urllib.request.urlopen(req, timeout=20) as r:
        rows = json.load(r)
    out = []
    for row in rows:
        tr = row.pop("track", None) or {}
        row["meta"] = tr.get("meta")          # 진단 요약(구형 기록엔 없음)
        row["elev"] = tr.get("elev")
        row["points"] = len(tr.get("points") or [])
        out.append(row)
    return out


# 등반 배터리 모드(IOS.md §7-5) — meta.bat_mode 원값 → 표시 라벨. admin.js 와 같은 표기.
BAT_MODE_LABEL = {"normal": "일반", "saver": "절전", "max": "최대절전"}


def _records_xlsx(limit=1000):
    """등반 기록 → xlsx (관리자 표와 같은 분리 열 + 분석용 수치 컬럼)."""
    import io
    from datetime import datetime, timedelta
    import openpyxl

    def kst(iso):
        if not iso:
            return None
        try:
            return (datetime.fromisoformat(iso.replace("Z", "+00:00"))
                    + timedelta(hours=9)).strftime("%Y-%m-%d %H:%M")
        except ValueError:
            return iso

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "records"
    # 저전력·전경/배경은 §7-4 계측 보강(meta v3) — 소모율만으로는 원인을 못 가른다.
    # 전경/배경은 분석용이라 분(minute) 수치로 따로 낸다(피벗·산점도에 바로 쓰도록).
    # 배터리 모드는 §7-5 — 모드별 %/h 비교의 축이라 피벗 첫 열로 쓰기 좋게 앞쪽에 둔다.
    head = ["시작(KST)", "종료(KST)", "산코드", "코스", "거리(km)", "시간(분)",
            "누적상승(m)", "배터리 모드",
            "배터리 시작(%)", "배터리 끝(%)", "충전중", "저전력", "소모율(%/h)",
            "전경(분)", "배경(분)",
            "GPS 수준", "GPS 지점", "GPS 유실", "평균정확도(m)",
            "빌드", "기기", "OS", "트랙 점수", "사용자", "기록 id"]
    ws.append(head)
    for r in _records(limit):
        m = r.get("meta") or {}
        rate = None
        if m and not m.get("charged") and m.get("bat_start", -1) >= 0 \
                and m.get("bat_end", -1) >= 0 and (r.get("duration_s") or 0) >= 600:
            rate = round((m["bat_start"] - m["bat_end"]) / (r["duration_s"] / 3600), 1)
        ws.append([
            kst(r.get("started_at")), kst(r.get("ended_at")),
            r.get("mountain_id"), r.get("course_name"),
            r.get("distance_km"),
            round(r["duration_s"] / 60) if r.get("duration_s") else None,
            r.get("ascent_m"),
            BAT_MODE_LABEL.get(m.get("bat_mode"), m.get("bat_mode")),   # v4~
            m.get("bat_start") if m.get("bat_start", -1) >= 0 else None,
            m.get("bat_end") if m.get("bat_end", -1) >= 0 else None,
            bool(m.get("charged")) if m else None,
            bool(m["lpm"]) if m.get("lpm") is not None else None,   # v3~
            rate,
            round(m["fg_s"] / 60) if m.get("fg_s") is not None else None,
            round(m["bg_s"] / 60) if m.get("bg_s") is not None else None,
            m.get("gps_mode"), m.get("fixes"), m.get("fixes_dropped"),
            m.get("acc_avg") if m.get("acc_avg", -1) >= 0 else None,
            f"b{m['build']}" if m.get("build") else None,
            m.get("device"), m.get("os"),
            r.get("points"), (r.get("user_id") or "")[:8], r.get("id"),
        ])
    ws.freeze_panes = "A2"
    for i in range(1, len(head) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 12
    ws.column_dimensions["A"].width = ws.column_dimensions["B"].width = 16
    ws.column_dimensions["D"].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class AdminHandler(BaseHandler):
    # ── 공통 ──
    def _json(self, obj, status=200):
        body = json.dumps(obj, ensure_ascii=False).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _err(self, msg, status=400):
        self._json({"error": msg}, status)

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        return self.rfile.read(n) if n else b""

    def _authorized(self, q):
        if self.client_address[0] in ("127.0.0.1", "::1"):
            return True
        tok = os.environ.get("ADMIN_TOKEN")
        if not tok:
            return False
        return (self.headers.get("X-Admin-Token") == tok
                or (q.get("token") or [None])[0] == tok)

    _KMA_OPS = {"ncst": "getUltraSrtNcst", "ufcst": "getUltraSrtFcst", "vfcst": "getVilageFcst"}

    def _weather(self, q):
        """기상청 단기예보 프록시 (CORS 회피 + 키 은닉). api/weather.js 와 동일 계약."""
        op = (q.get("op") or [""])[0]
        path = self._KMA_OPS.get(op)
        need = {k: (q.get(k) or [""])[0] for k in ("nx", "ny", "base_date", "base_time")}
        if not path or not all(need.values()):
            return self._err("bad params", 400)
        key = os.environ.get("KMA_KEY") or os.environ.get("KNPS_KEY")
        if not key:
            return self._err("no KMA_KEY", 500)
        sk = key if re.search(r"%[0-9A-Fa-f]{2}", key) else urllib.parse.quote(key, safe="")
        qs = urllib.parse.urlencode({"dataType": "JSON", "numOfRows": "300", "pageNo": "1", **need})
        url = (f"https://apis.data.go.kr/1360000/VilageFcstInfoService_2.0/{path}"
               f"?serviceKey={sk}&{qs}")
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "hiheight/1.0"})
            resp = urllib.request.urlopen(req, timeout=20)
            body = resp.read()
        except Exception as e:
            return self._err(f"upstream: {e}", 502)
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "public, max-age=600")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _login(self):
        """비밀번호(.env ADMIN_PASSWORD) 검증 → 성공 시 API 토큰 발급."""
        ip = self.client_address[0]
        now = time.time()
        rec = LOGIN_FAILS.setdefault(ip, {"n": 0, "until": 0.0})
        if now < rec["until"]:
            wait = int(rec["until"] - now)
            return self._json({"error": f"로그인 잠금 — {wait // 60 + 1}분 후 다시 시도하세요",
                               "retry_after": wait}, 429)
        try:
            data = json.loads(self._body() or b"{}")
        except ValueError:
            data = {}
        pw = os.environ.get("ADMIN_PASSWORD")
        if pw and str(data.get("password", "")) == pw:
            LOGIN_FAILS.pop(ip, None)
            return self._json({"token": os.environ["ADMIN_TOKEN"]})
        rec["n"] += 1
        if rec["n"] >= MAX_FAILS:
            rec["n"] = 0
            rec["until"] = now + LOCK_SEC
            return self._json({"error": f"{MAX_FAILS}회 실패 — 10분 동안 접속이 차단됩니다",
                               "retry_after": LOCK_SEC}, 429)
        return self._json({"error": f"비밀번호가 틀렸습니다 (남은 시도 {MAX_FAILS - rec['n']}회)"}, 401)

    def _route(self, method):
        u = urllib.parse.urlparse(self.path)
        q = urllib.parse.parse_qs(u.query)
        parts = [p for p in u.path.split("/") if p]  # ["api", ...]
        if method == "POST" and parts[1:] == ["login"]:
            return self._login()  # 로그인은 인증 없이 (실패 잠금으로 보호)
        if method == "GET" and parts[1:] == ["weather"]:
            return self._weather(q)  # 앱 공개 기능 — 인증 제외
        if not self._authorized(q):
            return self._err("관리자 인증 필요", 401)
        try:
            return self._dispatch(method, parts[1:], q)
        except FileNotFoundError as e:
            self._err(str(e), 404)
        except ValueError as e:
            self._err(str(e), 400)
        except Exception as e:
            self._err(f"{type(e).__name__}: {e}", 500)
            traceback.print_exc()

    # ── 라우팅 ──
    def _dispatch(self, method, p, q):
        # GET /api/version — 배포된 커밋(관리자 우상단 표시)
        if method == "GET" and p == ["version"]:
            return self._json(git_version())

        # GET /api/weather?op=&nx=&ny=&base_date=&base_time=  (기상청 프록시)
        if method == "GET" and p == ["weather"]:
            return self._weather(q)

        # GET /api/records?limit= — 등반 기록 + 진단 요약(배터리·GPS). service_role 조회.
        if method == "GET" and p == ["records"]:
            try:
                limit = int((q.get("limit") or ["200"])[0])
            except ValueError:
                limit = 200
            try:
                return self._json(_records(min(max(limit, 1), 1000)))
            except Exception as e:
                return self._err(str(e), 500)

        # GET /api/records/export?limit= — 등반 기록 xlsx 다운로드
        if method == "GET" and p == ["records", "export"]:
            try:
                limit = int((q.get("limit") or ["1000"])[0])
            except ValueError:
                limit = 1000
            data = _records_xlsx(min(max(limit, 1), 1000))
            self.send_response(200)
            self.send_header("Content-Type",
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition",
                             f'attachment; filename="hiheight-records-{time.strftime("%Y%m%d")}.xlsx"')
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(data)
            return

        # ── 시트 — 산 메타·스팟 일괄 편집 (그리드 JSON + xlsx 왕복, sheet_lib) ──
        # GET /api/sheet — 그리드 데이터 + 백업 목록
        if method == "GET" and p == ["sheet"]:
            return self._json({**sheet_lib.collect(), "backups": sheet_lib.list_backups()})

        # GET /api/sheet/export — 전 산 통합 xlsx (mountains/spots/설명 3시트)
        if method == "GET" and p == ["sheet", "export"]:
            data = sheet_lib.export_xlsx()
            self.send_response(200)
            self.send_header("Content-Type",
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition",
                             f'attachment; filename="hiheight-sheet-{time.strftime("%Y%m%d")}.xlsx"')
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(data)
            return

        # POST /api/sheet/import[?apply=1] — xlsx 업로드 → 변경 계획(diff) / 적용
        # apply 없이 호출하면 미리보기(무적용), apply=1 이면 스냅샷 백업 후 병합.
        if method == "POST" and p == ["sheet", "import"]:
            body = self._body()
            if not body:
                raise ValueError("xlsx 본문 없음")
            plan = sheet_lib.build_plan(sheet_lib.parse_xlsx(body),
                                        set(SPOT_DISPLAY_DEFAULT["categories"]))
            if plan["errors"] or (q.get("apply") or ["0"])[0] != "1":
                return self._json(sheet_lib.public_plan(plan))
            return self._json({**sheet_lib.apply_plan(plan), **sheet_lib.public_plan(plan)})

        # PUT /api/sheet — 그리드 저장 (행 스키마는 xlsx 와 동일, 즉시 적용 + 자동 백업)
        if method == "PUT" and p == ["sheet"]:
            data = json.loads(self._body())
            plan = sheet_lib.build_plan(data, set(SPOT_DISPLAY_DEFAULT["categories"]))
            if plan["errors"]:
                return self._json(sheet_lib.public_plan(plan), 400)
            return self._json({**sheet_lib.apply_plan(plan), **sheet_lib.public_plan(plan)})

        # POST /api/sheet/restore?ts= — 백업 스냅샷 복원
        if method == "POST" and p == ["sheet", "restore"]:
            return self._json(sheet_lib.restore((q.get("ts") or [""])[0]))

        # GET /api/mnt-codes?q=
        if method == "GET" and p == ["mnt-codes"]:
            kw = (q.get("q") or [""])[0].strip()
            rows = [m for m in _mnt_codes()
                    if kw and (kw in m["name"] or m["code"].startswith(kw))][:30]
            out = [dict(m, has_src=os.path.isdir(os.path.join(ROOT, "mountain", m["code"])),
                        has_draft=draft_store.load(m["code"]) is not None) for m in rows]
            return self._json(out)

        # GET/PUT /api/config/spots|poi — 표시 정책 (분류별 줌·기호·크기·볼드, 앱 전역)
        # 저장: admin_data/<이름>.json (원본) + R2 config/<이름>.json (앱이 부팅 시 fetch)
        # GET 도 정규화를 거쳐 레거시 파일(v1, 값=숫자)이 항상 v2 객체로 보인다.
        if len(p) == 2 and p[0] == "config" and p[1] in ("spots", "poi"):
            defaults = SPOT_DISPLAY_DEFAULT if p[1] == "spots" else POI_DISPLAY_DEFAULT
            fname = "spot-display.json" if p[1] == "spots" else "poi-display.json"
            cfg_path = os.path.join(draft_store.ADMIN_DATA, fname)
            if method == "GET":
                cats = {}
                if os.path.exists(cfg_path):
                    cats = json.load(open(cfg_path, encoding="utf-8")).get("categories", {})
                return self._json(_norm_display_cfg(cats, defaults))
            if method == "PUT":
                data = json.loads(self._body())
                cats = data.get("categories")
                if not isinstance(cats, dict):
                    raise ValueError("categories 객체 필요")
                cfg = _norm_display_cfg(cats, defaults)
                body = json.dumps(cfg, ensure_ascii=False, indent=1).encode()
                os.makedirs(draft_store.ADMIN_DATA, exist_ok=True)
                with open(cfg_path, "wb") as f:
                    f.write(body)
                import r2_lib
                r2_lib.upload_bytes(body, f"config/{fname}",
                                    content_type="application/json")
                return self._json({"ok": True, **cfg})

        # GET/PUT /api/config/curations — 큐레이션(추천 모음: 산·코스) 문서 전체 교체 방식
        # 저장: admin_data/curations.json (원본) + R2 config/curations.json (앱 추천 탭이 fetch)
        if p == ["config", "curations"]:
            cfg_path = os.path.join(draft_store.ADMIN_DATA, "curations.json")
            if method == "GET":
                if os.path.exists(cfg_path):
                    return self._json(json.load(open(cfg_path, encoding="utf-8")))
                return self._json({"version": 1, "curations": []})
            if method == "PUT":
                data = json.loads(self._body())
                cus = data.get("curations")
                if not isinstance(cus, list):
                    raise ValueError("curations 배열 필요")
                for cu in cus:
                    if not isinstance(cu, dict) or not str(cu.get("title", "")).strip():
                        raise ValueError("큐레이션 이름(title)이 비어 있음")
                    if not isinstance(cu.get("items"), list):
                        raise ValueError(f"{cu['title']}: items 배열 필요")
                    for it in cu["items"]:
                        if it.get("type") not in ("mountain", "course"):
                            raise ValueError(f"{cu['title']}: 항목 type 은 mountain|course")
                        if not re.fullmatch(r"\d{9}", str(it.get("code", ""))):
                            raise ValueError(f"{cu['title']}: 항목 code 는 9자리 산코드")
                        if not str(it.get("name", "")).strip():
                            raise ValueError(f"{cu['title']}: 항목 name 필요")
                        if it.get("img") is not None and not str(it["img"]).startswith("https://"):
                            raise ValueError(f"{cu['title']}: img 는 https URL")
                cfg = {"version": 1, "curations": [
                    {"id": cu.get("id") or f"cu-{uuid.uuid4().hex[:8]}",
                     "title": str(cu["title"]).strip(),
                     # 슬라이드 요소(전부 선택 — 비면 미표시): sub 부가설명(반투명 배지),
                     # title 큰 제목, desc 중앙 하단 설명, logo 좌하단 마크,
                     # credit 우하단 출처(사진 저작자), img 배경 이미지
                     "items": [{k: it[k] for k in
                                ("type", "code", "name", "mountain",
                                 "sub", "title", "desc", "logo", "credit", "img")
                                if it.get(k) not in (None, "")} for it in cu["items"]]}
                    for cu in cus]}
                body = json.dumps(cfg, ensure_ascii=False, indent=1).encode()
                os.makedirs(draft_store.ADMIN_DATA, exist_ok=True)
                with open(cfg_path, "wb") as f:
                    f.write(body)
                import r2_lib
                # Cache-Control 은 r2_lib.cache_control_for 가 확장자로 결정한다
                # (*.json → no-cache). 없으면 앱이 옛 큐레이션을 몇 시간 재사용한다.
                r2_lib.upload_bytes(body, "config/curations.json",
                                    content_type="application/json")
                return self._json({"ok": True, **cfg})

        # POST /api/mountain-image?code= — 산 커버 이미지 업로드 (큐레이션 캐러셀 배경)
        # 본문 = 이미지 바이트 그대로 (Content-Type 으로 형식 판별) → R2 images/mountains/
        if method == "POST" and p == ["mountain-image"]:
            code = (q.get("code") or [""])[0]
            if not re.fullmatch(r"\d{9}", code):
                raise ValueError("9자리 산코드 필요")
            ctype = (self.headers.get("Content-Type") or "").split(";")[0].strip()
            ext = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp"}.get(ctype)
            if not ext:
                raise ValueError("이미지 형식은 jpeg/png/webp")
            body = self._body()
            if not body:
                raise ValueError("이미지 본문 없음")
            if len(body) > 8 * 1024 * 1024:
                raise ValueError("이미지는 8MB 이하")
            import r2_lib
            key = f"images/mountains/{code}.{ext}"
            r2_lib.upload_bytes(body, key, content_type=ctype)
            return self._json({"url": f"{r2_lib.public_base()}/{key}"})

        # GET /api/course-search?q= — 전체 초안의 코스를 이름으로 검색 (큐레이션 항목 추가용)
        if method == "GET" and p == ["course-search"]:
            kw = (q.get("q") or [""])[0].strip()
            out = []
            if kw:
                for m in draft_store.list_drafts():
                    d = draft_store.load(m["code"])
                    for c in d["courses"]:
                        name = c.get("name") or ""
                        if kw in name or kw in m["name"]:
                            out.append({"code": m["code"], "mountain": m["name"],
                                        "name": name, "no": c.get("no"),
                                        "status": c["status"],
                                        "published": m.get("published", True)})
                        if len(out) >= 30:
                            break
                    if len(out) >= 30:
                        break
            return self._json(out)

        # GET/POST /api/mountains
        if p == ["mountains"]:
            if method == "GET":
                return self._json(draft_store.list_drafts())
            if method == "POST":
                data = json.loads(self._body() or b"{}")
                code = data.get("code", "")
                if not re.fullmatch(r"\d{9}", code):
                    raise ValueError("9자리 산코드 필요")
                draft, jid = create_draft(code)
                return self._json({"code": code, "job_id": jid}, 201)

        # POST /api/mountains/manual — 산림청 원본 없는 산 직접 등록 (이름·중심·높이)
        # 코스 자동 시드/DEM 프리페치 없음(코스는 GPX 업로드) → 즉시 초안만 생성.
        if method == "POST" and p == ["mountains", "manual"]:
            data = json.loads(self._body() or b"{}")
            name = str(data.get("name", "")).strip()
            if not name:
                raise ValueError("산 이름이 필요합니다")
            try:
                lat, lon = float(data["lat"]), float(data["lon"])
            except (KeyError, TypeError, ValueError):
                raise ValueError("중심 좌표(lat·lon)가 필요합니다")
            if not (32 <= lat <= 44 and 123 <= lon <= 133):
                raise ValueError("좌표가 한반도 범위(위도 32~44, 경도 123~133)를 벗어났습니다")
            elev = data.get("elev")
            if elev in (None, ""):
                elev = None
            else:
                try:
                    elev = int(float(elev))
                except (TypeError, ValueError):
                    raise ValueError("높이는 숫자여야 합니다")
            region = str(data.get("region", "")).strip() or None
            code = _next_manual_code()
            if draft_store.load(code):
                raise ValueError(f"{code} 초안이 이미 있음")
            draft = draft_store.new_manual_draft(code, name, [lon, lat], elev, region)
            draft_store.save(code, draft)
            return self._json({"code": code, "name": name}, 201)

        # /api/mountains/<code>/...
        if len(p) >= 2 and p[0] == "mountains" and re.fullmatch(r"\d{9}", p[1]):
            code, rest = p[1], p[2:]
            if not rest and method == "DELETE":
                # 앱에서 완전 삭제: Supabase 카탈로그 행 + R2 팩 파일 + 로컬 초안.
                # (초안이 이미 없어도 code 로 배포본을 정리 — 고아 배포본 대응)
                if JOBS.busy(code):
                    raise ValueError("이 산의 배포/준비 작업이 진행 중입니다 — 완료 후 삭제하세요"
                                     " (진행 중 삭제하면 고아 배포본이 생깁니다)")
                import shutil
                import publish_pack
                result = {"code": code}
                try:
                    result.update(publish_pack.unpublish(code))  # 행 + R2 파일
                except Exception as e:  # 네트워크/권한 실패는 보고하되 초안 제거는 진행
                    result["unpublish_error"] = str(e)
                d = os.path.join(draft_store.ADMIN_DATA, code)
                if os.path.isdir(d):
                    shutil.rmtree(d)
                    result["draft_removed"] = True
                return self._json({"ok": True, **result})
            if rest == ["draft"]:
                if method == "GET":
                    d = draft_store.load(code)
                    if not d:
                        raise FileNotFoundError(f"{code} 초안 없음")
                    return self._json(d)
                if method == "PUT":
                    # 갱신 전용 — 생성은 POST /api/mountains 만. 삭제 직후/다른 탭의
                    # 늦은 자동저장이 지워진 초안을 되살리는 것을 차단.
                    if not draft_store.load(code):
                        raise FileNotFoundError(f"{code} 초안 없음 (삭제되었거나 미등록)")
                    d = json.loads(self._body())
                    if d.get("mountain", {}).get("code") != code:
                        raise ValueError("mountain.code 불일치")
                    draft_store.save(code, d)
                    return self._json({"ok": True})
            if rest == ["network"] and method == "GET":
                segs, _ = pl.load_forest_segments(code)
                return self._json(pl.network_geojson(segs))
            if rest == ["contours"] and method == "GET":
                # 관리자 지도 등고선 — 배포 산출물(data/packs) 재사용, 없으면 DEM 즉석 생성 후 캐시.
                # 전국 DEM 프리페치 완료 상태라 생성도 로컬에서 몇 초면 끝난다.
                d = draft_store.load(code)
                if not d:
                    raise FileNotFoundError(f"{code} 초안 없음")
                pack_f = os.path.join(ROOT, "data", "packs", code, "contours.geojson")
                cache_f = os.path.join(ROOT, "cache", "contours", f"{code}.geojson")
                src = pack_f if os.path.exists(pack_f) else cache_f
                if not os.path.exists(src):
                    import dem_cache
                    import make_contours_copernicus as mc
                    bbox = d["mountain"]["bbox"]
                    os.makedirs(os.path.dirname(cache_f), exist_ok=True)
                    mc.generate(cache_f, bbox, dem_cache.ensure(bbox), log=lambda s: None)
                    src = cache_f
                return self._json(json.load(open(src, encoding="utf-8")))
            if rest == ["gpx"] and method == "POST":
                import gpx_match
                d = draft_store.load(code)
                if not d:
                    raise FileNotFoundError(f"{code} 초안 없음")
                tau = float((q.get("tau") or [25])[0])
                detour = float((q.get("detour") or [1.6])[0])
                upload_name = (q.get("name") or [""])[0]  # 원본 파일명 — 포맷 판별·보존용
                raw = self._body()
                return self._json(gpx_match.match_gpx_upload(code, d, raw, tau, detour, upload_name))
            if len(rest) == 3 and rest[0] == "courses" and rest[2] == "recompute" \
                    and method == "POST":
                import gpx_match
                d = draft_store.load(code)
                cid = rest[1]
                c = next((c for c in d["courses"] if c["id"] == cid), None)
                if not c:
                    raise FileNotFoundError(f"코스 {cid} 없음")
                gpx_match.recompute_course(d, c)
                draft_store.save(code, d)
                return self._json(c)
            if rest == ["publish"] and method == "POST":
                import publish_pack
                d = draft_store.load(code)
                if not d:
                    raise FileNotFoundError(f"{code} 초안 없음")
                jid = JOBS.submit(f"{d['mountain']['name']} 배포",
                                  lambda job: publish_pack.publish(code, job, job_step),
                                  code=code)
                return self._json({"job_id": jid})


        # GET /api/jobs/<id>
        if method == "GET" and len(p) == 2 and p[0] == "jobs":
            job = JOBS.jobs.get(p[1])
            if not job:
                raise FileNotFoundError("잡 없음")
            return self._json(job)

        self._err("no route", 404)

    # ── HTTP 메서드 ──
    def do_GET(self):
        if self.path.startswith("/api/"):
            return self._route("GET")
        return super().do_GET()

    def do_POST(self):
        if self.path.startswith("/api/"):
            return self._route("POST")
        self.send_error(405)

    def do_PUT(self):
        if self.path.startswith("/api/"):
            return self._route("PUT")
        self.send_error(405)

    def do_DELETE(self):
        if self.path.startswith("/api/"):
            return self._route("DELETE")
        self.send_error(405)

    def log_message(self, fmt, *args):  # /api 폴링 소음 축소
        # args[0] 이 HTTPStatus 등 비문자열일 수 있음(send_error 경로) — str 로 안전 변환
        if "/api/jobs/" not in str(args[0] if args else ""):
            super().log_message(fmt, *args)


def main():
    load_env()
    tok = ensure_admin_token()
    # (주의) 레거시 역임포트는 완전 폐기 — 시작 시 자동 임포트는 삭제한 산을 재시작마다
    # 되살렸고, 등록 시 LEGACY 분기는 거친 KNPS 선형(점 간격 ~104m)을 다시 들여왔다.
    # 모든 산은 산림청 원본 기반 신규 등록 경로만 사용한다.
    handler = partial(AdminHandler, directory=ROOT)
    httpd = ThreadingHTTPServer((BIND, PORT), handler)
    print(f"관리자 콘솔: http://{BIND}:{PORT}/admin/?token={tok}")
    print("(localhost 접속은 토큰 불필요 · 사용자 앱도 같은 포트 / 에서 서빙)")
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        httpd.shutdown()


if __name__ == "__main__":
    main()
