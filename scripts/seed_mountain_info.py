#!/usr/bin/env python3
"""산림청 산정보(scripts/mnt.xlsx) + 산코드 목록(scripts/MNT_CODE.xlsx) 병합
→ Supabase mountain_info 테이블 시드 (data/mnt-codes.json 과 동일한 5,360건 기준).

mountain_info 는 전국 산 카탈로그(공개 읽기):
  code(산코드 PK) · name · region(소재지) · elev(높이 m) ·
  manager(관리주체) · manager_tel(관리자 전화) · description(산 설명) · data_date(기준일)

사전조건: supabase/schema.sql 의 mountain_info 블록이 실행되어 있어야 한다.

사용: (venv 에 openpyxl 필요)
  export SUPABASE_URL=https://durnojryhhsajnlwvdzt.supabase.co
  export SUPABASE_SECRET_KEY=sb_secret_...   # ⚠️ 비밀! .env 로만, 커밋 금지
  python3 scripts/seed_mountain_info.py
"""
import json
import os
import sys
import urllib.error
import urllib.request

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl 필요: pip install openpyxl")

URL = os.environ.get("SUPABASE_URL", "").rstrip("/")
KEY = os.environ.get("SUPABASE_SECRET_KEY") or os.environ.get("SUPABASE_SERVICE_ROLE", "")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "scripts", "mnt.xlsx")

if not URL or not KEY:
    sys.exit("환경변수 SUPABASE_URL, SUPABASE_SECRET_KEY 를 설정하세요.")

norm = lambda v: " ".join(str(v or "").split()) or None

wb = openpyxl.load_workbook(SRC, read_only=True)
ws = wb[wb.sheetnames[0]]
rows = []
# 열: 0 산코드, 1 산명, 6 소재지, 7 관리주체명, 8 관리자전화, 9 설명, 11 높이, 12 기준일자
for r in list(ws.iter_rows(values_only=True))[1:]:
    if r[0] is None or not r[1]:
        continue
    try:
        elev = float(r[11]) if r[11] is not None else None
    except (TypeError, ValueError):
        elev = None
    desc = norm(r[9])
    if desc in ("( - )", "-"):
        desc = None
    rows.append({
        "code": str(r[0]).strip(),
        "name": norm(r[1]),
        "region": norm(r[6]),
        "elev": elev if (elev or 0) > 0 else None,
        "manager": norm(r[7]),
        "manager_tel": norm(r[8]),
        "description": desc,
        "data_date": r[12].date().isoformat() if r[12] else None,
    })

# 병합: MNT_CODE.xlsx(목록)에만 있는 코드도 포함해 DB 를 전국 카탈로그(5,360건)와 일치시킴
# (이름/소재지만 보유 — 높이·관리주체·설명은 산정보 원본에 없는 코드라 null)
codes_seen = {r["code"] for r in rows}
SRC_LIST = os.path.join(ROOT, "scripts", "MNT_CODE.xlsx")
wb2 = openpyxl.load_workbook(SRC_LIST, read_only=True)
extra = 0
for _, name, region, code in list(wb2[wb2.sheetnames[0]].iter_rows(values_only=True))[1:]:
    if not code or not name:
        continue
    code = str(code).strip()
    if code in codes_seen:
        continue
    codes_seen.add(code)
    rows.append({"code": code, "name": norm(name), "region": norm(region),
                 "elev": None, "manager": None, "manager_tel": None,
                 "description": None, "data_date": None})
    extra += 1

print(f"산정보 {len(rows) - extra}건 + 목록 전용 {extra}건 병합 — 업서트 시작")

BATCH = 500
ok = fail = 0
for i in range(0, len(rows), BATCH):
    chunk = rows[i:i + BATCH]
    body = json.dumps(chunk).encode()
    req = urllib.request.Request(
        URL + "/rest/v1/mountain_info", data=body, method="POST",
        headers={"Authorization": "Bearer " + KEY, "apikey": KEY,
                 "Content-Type": "application/json",
                 "Prefer": "resolution=merge-duplicates,return=minimal"})
    try:
        resp = urllib.request.urlopen(req, timeout=120)
        ok += len(chunk)
        print(f"  {i + len(chunk)}/{len(rows)} upsert OK ({resp.status})")
    except urllib.error.HTTPError as e:
        fail += len(chunk)
        print(f"  {i}~ 실패({e.code}): {e.read()[:200]}")

print(f"완료: 성공 {ok} / 실패 {fail}")
